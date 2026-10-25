"""Indexa a pasta 'Maintenance Asset Management' em SQLite FTS5 (BM25). 100% local."""
import re, sqlite3, sys
from pathlib import Path

import fitz  # pymupdf
import openpyxl
import xlrd
from docx import Document
from pptx import Presentation

HERE = Path(__file__).parent
SRC = HERE.parent / "CMMS" / "Maintenance Asset Management"
DB = HERE / "rag.db"
CHUNK, OVERLAP = 1800, 250


def pdf(p):
    with fitz.open(p) as d:
        for i, page in enumerate(d, 1):
            yield f"p.{i}", page.get_text("text")


def docx(p):
    d = Document(p)
    parts = [x.text for x in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            parts.append(" | ".join(c.text.strip() for c in r.cells))
    yield "", "\n".join(parts)


def pptx(p):
    for i, s in enumerate(Presentation(p).slides, 1):
        txt = []
        for sh in s.shapes:
            if sh.has_text_frame:
                txt.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                for r in sh.table.rows:
                    txt.append(" | ".join(c.text for c in r.cells))
        if s.has_notes_slide:
            txt.append(s.notes_slide.notes_text_frame.text)
        yield f"slide {i}", "\n".join(txt)


def xlsx(p):
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = [" | ".join(str(c) for c in r if c is not None) for r in ws.iter_rows(values_only=True)]
        yield f"sheet {ws.title}", "\n".join(r for r in rows if r.strip())


def xls(p):
    try:
        wb = xlrd.open_workbook(p)
    except Exception:
        # Alguns .xls exportados são HTML/XML disfarçados
        raw = Path(p).read_bytes().decode("utf-8", "ignore")
        yield "", re.sub(r"<[^>]+>", " ", raw)
        return
    for sh in wb.sheets():
        rows = [" | ".join(str(c.value) for c in sh.row(r) if str(c.value).strip()) for r in range(sh.nrows)]
        yield f"sheet {sh.name}", "\n".join(r for r in rows if r.strip())


def plain(p):
    yield "", Path(p).read_text("utf-8", "ignore")


READERS = {".pdf": pdf, ".docx": docx, ".pptx": pptx, ".xlsx": xlsx, ".xls": xls,
           ".txt": plain, ".md": plain, ".textclipping": plain}


def chunks(text):
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    i = 0
    while i < len(text):
        end = min(i + CHUNK, len(text))
        if end < len(text):
            cut = text.rfind("\n", i + CHUNK // 2, end)
            end = cut if cut > 0 else end
        piece = text[i:end].strip()
        if len(piece) > 40:
            yield piece
        if end >= len(text):
            break
        i = max(end - OVERLAP, i + 1)


def main():
    DB.unlink(missing_ok=True)
    con = sqlite3.connect(DB)
    con.execute("CREATE VIRTUAL TABLE chunks USING fts5(content, path UNINDEXED, loc UNINDEXED, "
                "area, tokenize='porter unicode61 remove_diacritics 2')")
    n_files = n_chunks = 0
    for f in sorted(SRC.rglob("*")):
        if not f.is_file() or f.suffix.lower() not in READERS:
            continue
        rel = f.relative_to(SRC).as_posix()
        area = " ".join(rel.split("/")[:-1])
        try:
            rows = []
            for loc, text in READERS[f.suffix.lower()](f):
                for c in chunks(text):
                    rows.append((f"{f.stem}\n{c}", rel, loc, area))
            con.executemany("INSERT INTO chunks VALUES (?,?,?,?)", rows)
            con.commit()
            n_files += 1
            n_chunks += len(rows)
            print(f"{len(rows):5d}  {rel}", flush=True)
        except Exception as e:
            print(f"FALHA {rel}: {e}", file=sys.stderr)
    con.execute("INSERT INTO chunks(chunks) VALUES('optimize')")
    con.commit()
    print(f"\n{n_files} arquivos, {n_chunks} chunks -> {DB}")


if __name__ == "__main__":
    main()
